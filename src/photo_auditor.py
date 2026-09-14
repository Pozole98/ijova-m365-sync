"""
Módulo de Auditoría y Descarga de Fotografías de Perfil de Alumnos (Microsoft 365 / Entra ID).
Permite descargar masivamente las fotos de perfil de los alumnos, identificar quiénes tienen
fotografía configurada y generar una galería web interactiva para verificar si son apropiadas.
"""
import os
import re
import csv
import json
import base64
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor, as_completed
from typing import List, Dict, Any, Optional, Tuple
from pydantic import BaseModel, Field

from src.config import AppConfig
from src.graph_client import GraphClient, GraphClientError
from src.models import StudentRecord, EntraUser


class PhotoAuditRecord(BaseModel):
    matricula: str
    upn: str
    display_name: str
    nivel: str = "No especificado"
    grado_semestre: str = ""
    has_photo: bool = False
    photo_filename: Optional[str] = None
    photo_rel_path: Optional[str] = None
    photo_size_bytes: int = 0
    photo_base64: Optional[str] = None
    content_type: str = "image/jpeg"
    audit_date: str = Field(default_factory=lambda: datetime.now().strftime("%Y-%m-%d %H:%M:%S"))


def _slugify(text: str) -> str:
    """Convierte texto en un slug seguro para nombres de archivo."""
    text = text.upper().strip()
    text = re.sub(r"[ÁÀÄÂ]", "A", text)
    text = re.sub(r"[ÉÈËÊ]", "E", text)
    text = re.sub(r"[ÍÌÏÎ]", "I", text)
    text = re.sub(r"[ÓÒÖÔ]", "O", text)
    text = re.sub(r"[ÚÙÜÛ]", "U", text)
    text = re.sub(r"[Ñ]", "N", text)
    text = re.sub(r"[^A-Z0-9_-]", "_", text)
    text = re.sub(r"_+", "_", text).strip("_")
    return text[:40]


def build_student_lookup(config: AppConfig) -> Dict[str, Dict[str, str]]:
    """
    Construye un mapa de consulta a partir del archivo Excel/ODS matriculado
    para enriquecer los datos de Entra ID con Nivel y Grado.
    """
    lookup: Dict[str, Dict[str, str]] = {}
    excel_path = config.excel_path
    if not os.path.exists(excel_path):
        return lookup

    try:
        from src.excel_parser import parse_excel_students
        students = parse_excel_students(excel_path, config.sheet_name)
        for s in students:
            mat = s.matricula.strip().lower()
            upn = s.upn_raw.strip().lower()
            info = {
                "matricula": s.matricula,
                "display_name": f"{s.apellido_paterno} {s.apellido_materno} {s.nombres}".strip(),
                "nivel": s.nivel or "No especificado",
                "grado": s.grado_semestre or ""
            }
            if mat:
                lookup[mat] = info
            if upn:
                lookup[upn] = info
    except Exception as e:
        print(f"⚠️ Nota: No se pudo cruzar con el Excel ({e}). Se usarán datos directos de Entra ID.")

    return lookup


def download_single_photo(
    graph: GraphClient,
    user: EntraUser,
    student_info: Dict[str, str],
    photos_dir: str
) -> PhotoAuditRecord:
    """
    Consulta y descarga la fotografía de un usuario individual.
    """
    matricula = (user.mail_nickname or user.user_principal_name.split("@")[0]).strip()
    display_name = user.display_name or student_info.get("display_name", matricula)
    nivel = student_info.get("nivel", "No especificado")
    grado = student_info.get("grado", "")
    upn = user.user_principal_name.strip().lower()

    photo_bytes = graph.get_user_photo(user.id or upn)

    if photo_bytes:
        slug = _slugify(display_name)
        filename = f"{matricula}_{slug}.jpg"
        file_path = os.path.join(photos_dir, filename)
        
        try:
            with open(file_path, "wb") as f:
                f.write(photo_bytes)
        except Exception:
            pass

        b64_img = base64.b64encode(photo_bytes).decode("utf-8")
        rel_path = f"fotos_perfil/{filename}"

        return PhotoAuditRecord(
            matricula=matricula,
            upn=upn,
            display_name=display_name,
            nivel=nivel,
            grado_semestre=grado,
            has_photo=True,
            photo_filename=filename,
            photo_rel_path=rel_path,
            photo_size_bytes=len(photo_bytes),
            photo_base64=b64_img,
            content_type="image/jpeg"
        )
    else:
        return PhotoAuditRecord(
            matricula=matricula,
            upn=upn,
            display_name=display_name,
            nivel=nivel,
            grado_semestre=grado,
            has_photo=False,
            photo_size_bytes=0
        )


def audit_profile_photos(
    graph: GraphClient,
    config: AppConfig,
    max_workers: int = 8,
    output_dir: Optional[str] = None
) -> Tuple[List[PhotoAuditRecord], Dict[str, Any], str, str]:
    """
    Flujo maestro de auditoría de fotografías de perfil:
    1. Obtiene usuarios de Entra ID.
    2. Cruza con listado escolar.
    3. Descarga fotografías concurrentemente.
    4. Genera Galería Web HTML y reporte Excel/CSV.
    """
    base_out = output_dir or config.reports_dir
    os.makedirs(base_out, exist_ok=True)
    photos_dir = os.path.join(base_out, "fotos_perfil")
    os.makedirs(photos_dir, exist_ok=True)

    print("\n" + "=" * 80)
    print("🖼️ AUDITORÍA Y RECUPERACIÓN DE FOTOS DE PERFIL — MICROSOFT 365")
    print("=" * 80)
    print(f"📁 Directorio de resguardo de imágenes: {photos_dir}")

    # 1. Cruzar datos con Excel si está disponible
    student_lookup = build_student_lookup(config)

    # 2. Obtener usuarios de Entra ID
    print("\n🔍 Consultando catálogo de usuarios en Microsoft Entra ID...")
    entra_users = graph.get_all_users()

    # Filtrar cuentas de alumnos (generalmente las que tienen matrícula o dominio ijova.com)
    # Excluir cuentas de servicio o administradores conocidos si no son alumnos
    target_users: List[EntraUser] = []
    admin_upn = (graph.admin_upn or "").lower()

    for u in entra_users:
        u_upn = u.user_principal_name.lower()
        if admin_upn and u_upn == admin_upn:
            continue
        # Incluir usuarios que pertenezcan al dominio configurado
        if config.domain.lower() in u_upn:
            target_users.append(u)

    print(f"\n⚡ Total de alumnos/usuarios a auditar: {len(target_users)}")
    print(f"🚀 Iniciando escaneo concurrente ({max_workers} hilos de descarga)...\n")

    records: List[PhotoAuditRecord] = []
    completed = 0
    total = len(target_users)
    photos_found = 0

    with ThreadPoolExecutor(max_workers=max_workers) as executor:
        future_to_user = {}
        for user in target_users:
            matricula = (user.mail_nickname or user.user_principal_name.split("@")[0]).strip().lower()
            info = student_lookup.get(matricula) or student_lookup.get(user.user_principal_name.lower(), {})
            future = executor.submit(download_single_photo, graph, user, info, photos_dir)
            future_to_user[future] = user

        for future in as_completed(future_to_user):
            completed += 1
            try:
                rec = future.result()
                records.append(rec)
                if rec.has_photo:
                    photos_found += 1
                    status_icon = "📸 \033[1;32mFOTO DETECTADA\033[0m"
                    size_kb = round(rec.photo_size_bytes / 1024, 1)
                    print(f"   [{completed:>3}/{total}] {status_icon} | {rec.matricula} - {rec.display_name} ({size_kb} KB)")
                else:
                    status_icon = "⚪ Sin foto"
                    print(f"   [{completed:>3}/{total}] {status_icon}       | {rec.matricula} - {rec.display_name}")
            except Exception as e:
                user = future_to_user[future]
                print(f"   [{completed:>3}/{total}] ❌ Error en {user.user_principal_name}: {e}")

    # Ordenar por: primero los que tienen foto, luego por matrícula
    records.sort(key=lambda r: (not r.has_photo, r.matricula))

    stats = {
        "total_audited": total,
        "with_photo": photos_found,
        "without_photo": total - photos_found,
        "coverage_pct": round((photos_found / total * 100) if total > 0 else 0, 1),
        "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    }

    print("\n" + "=" * 80)
    print("📊 RESUMEN DE AUDITORÍA DE FOTOGRAFÍAS:")
    print("=" * 80)
    print(f"   • Total de alumnos evaluados:       {stats['total_audited']}")
    print(f"   • Alumnos CON foto de perfil:      \033[1;32m{stats['with_photo']}\033[0m ({stats['coverage_pct']}%)")
    print(f"   • Alumnos SIN foto de perfil:      \033[1;33m{stats['without_photo']}\033[0m")
    print("=" * 80)

    # 3. Generar Galería Web Interactiva
    gallery_path = os.path.join(base_out, "galeria_fotos_alumnos.html")
    generate_html_gallery(records, gallery_path, stats, config)
    print(f"\n🌐 Galería Web Interactiva lista en: \033[1;34m{gallery_path}\033[0m")

    # 4. Generar Reporte Excel / CSV
    excel_path = os.path.join(base_out, "auditoria_fotos_perfil.xlsx")
    generate_excel_audit_report(records, excel_path, stats)
    csv_path = os.path.join(base_out, "auditoria_fotos_perfil.csv")
    generate_csv_audit_report(records, csv_path)
    print(f"📊 Reporte de Auditoría Excel en:   \033[1;32m{excel_path}\033[0m")

    return records, stats, gallery_path, excel_path


def generate_html_gallery(
    records: List[PhotoAuditRecord],
    output_path: str,
    stats: Dict[str, Any],
    config: Optional[AppConfig] = None
) -> str:
    """
    Genera una galería HTML interactiva, visualmente impecable con Vanilla CSS y JS.
    Permite auditar visualmente cada foto, marcar fotos inapropiadas y exportar listas.
    """
    domain = config.domain if config else "ijova.com"
    cards_json = []

    for r in records:
        cards_json.append({
            "matricula": r.matricula,
            "upn": r.upn,
            "name": r.display_name,
            "nivel": r.nivel,
            "grado": r.grado_semestre,
            "has_photo": r.has_photo,
            "size_kb": round(r.photo_size_bytes / 1024, 1),
            "filename": r.photo_filename or "",
            "img_src": f"data:image/jpeg;base64,{r.photo_base64}" if r.photo_base64 else "",
            "rel_path": r.photo_rel_path or ""
        })

    html_content = f"""<!DOCTYPE html>
<html lang="es">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>Auditoría de Fotos de Perfil — IJOVA</title>
  <style>
    :root {{
      --bg-primary: #0f172a;
      --bg-secondary: #1e293b;
      --bg-card: rgba(30, 41, 59, 0.7);
      --border-color: rgba(255, 255, 255, 0.1);
      --text-main: #f8fafc;
      --text-muted: #94a3b8;
      --accent-blue: #38bdf8;
      --accent-indigo: #6366f1;
      --accent-green: #22c55e;
      --accent-amber: #f59e0b;
      --accent-red: #ef4444;
      --card-radius: 16px;
      --transition: all 0.25s cubic-bezier(0.4, 0, 0.2, 1);
    }}

    * {{
      box-sizing: border-box;
      margin: 0;
      padding: 0;
    }}

    body {{
      font-family: 'Segoe UI', system-ui, -apple-system, Roboto, Helvetica, Arial, sans-serif;
      background: radial-gradient(circle at 10% 20%, #0f172a 0%, #020617 90%);
      color: var(--text-main);
      min-height: 100vh;
      padding: 2rem 1.5rem;
    }}

    .container {{
      max-width: 1400px;
      margin: 0 auto;
    }}

    /* Header */
    header {{
      background: linear-gradient(135deg, rgba(30, 41, 59, 0.8), rgba(15, 23, 42, 0.9));
      border: 1px solid var(--border-color);
      backdrop-filter: blur(12px);
      border-radius: var(--card-radius);
      padding: 2rem;
      margin-bottom: 2rem;
      box-shadow: 0 10px 30px -10px rgba(0, 0, 0, 0.5);
    }}

    .institution-badge {{
      display: inline-block;
      background: rgba(56, 189, 248, 0.15);
      color: var(--accent-blue);
      border: 1px solid rgba(56, 189, 248, 0.3);
      padding: 0.35rem 0.85rem;
      border-radius: 999px;
      font-size: 0.8rem;
      font-weight: 600;
      text-transform: uppercase;
      letter-spacing: 0.05em;
      margin-bottom: 0.75rem;
    }}

    h1 {{
      font-size: 2rem;
      font-weight: 800;
      letter-spacing: -0.02em;
      background: linear-gradient(to right, #ffffff, #94a3b8);
      -webkit-background-clip: text;
      -webkit-text-fill-color: transparent;
      margin-bottom: 0.5rem;
    }}

    .subtitle {{
      color: var(--text-muted);
      font-size: 0.95rem;
      display: flex;
      gap: 1.5rem;
      flex-wrap: wrap;
    }}

    /* Metrics Grid */
    .metrics-grid {{
      display: grid;
      grid-template-columns: repeat(auto-fit, minmax(220px, 1fr));
      gap: 1.25rem;
      margin-top: 1.75rem;
    }}

    .metric-card {{
      background: var(--bg-card);
      border: 1px solid var(--border-color);
      border-radius: 12px;
      padding: 1.25rem;
      position: relative;
      overflow: hidden;
      transition: var(--transition);
    }}

    .metric-card:hover {{
      transform: translateY(-2px);
      border-color: rgba(255, 255, 255, 0.2);
    }}

    .metric-card::before {{
      content: '';
      position: absolute;
      top: 0;
      left: 0;
      right: 0;
      height: 3px;
      background: var(--accent-indigo);
    }}

    .metric-card.green::before {{ background: var(--accent-green); }}
    .metric-card.amber::before {{ background: var(--accent-amber); }}
    .metric-card.red::before {{ background: var(--accent-red); }}

    .metric-value {{
      font-size: 2.2rem;
      font-weight: 800;
      line-height: 1;
      margin-top: 0.5rem;
    }}

    .metric-label {{
      font-size: 0.85rem;
      color: var(--text-muted);
      font-weight: 500;
    }}

    /* Toolbar */
    .toolbar {{
      background: var(--bg-secondary);
      border: 1px solid var(--border-color);
      border-radius: var(--card-radius);
      padding: 1.25rem;
      margin-bottom: 2rem;
      display: flex;
      flex-wrap: wrap;
      gap: 1rem;
      align-items: center;
      justify-content: space-between;
    }}

    .search-box {{
      flex: 1;
      min-width: 260px;
      position: relative;
    }}

    .search-box input {{
      width: 100%;
      background: rgba(15, 23, 42, 0.6);
      border: 1px solid var(--border-color);
      border-radius: 8px;
      padding: 0.75rem 1rem 0.75rem 2.5rem;
      color: #fff;
      font-size: 0.95rem;
      outline: none;
      transition: var(--transition);
    }}

    .search-box input:focus {{
      border-color: var(--accent-blue);
      box-shadow: 0 0 0 3px rgba(56, 189, 248, 0.2);
    }}

    .search-icon {{
      position: absolute;
      left: 0.85rem;
      top: 50%;
      transform: translateY(-50%);
      color: var(--text-muted);
      pointer-events: none;
    }}

    .filters {{
      display: flex;
      flex-wrap: wrap;
      gap: 0.5rem;
    }}

    .btn {{
      background: rgba(255, 255, 255, 0.05);
      border: 1px solid var(--border-color);
      color: var(--text-main);
      padding: 0.6rem 1rem;
      border-radius: 8px;
      font-size: 0.85rem;
      font-weight: 600;
      cursor: pointer;
      transition: var(--transition);
      display: inline-flex;
      align-items: center;
      gap: 0.4rem;
    }}

    .btn:hover {{
      background: rgba(255, 255, 255, 0.12);
      transform: translateY(-1px);
    }}

    .btn.active {{
      background: var(--accent-blue);
      color: #04101e;
      border-color: var(--accent-blue);
    }}

    .btn-danger {{
      background: rgba(239, 68, 68, 0.15);
      border-color: rgba(239, 68, 68, 0.4);
      color: #fca5a5;
    }}

    .btn-danger:hover {{
      background: var(--accent-red);
      color: #fff;
    }}

    /* Gallery Grid */
    .gallery-grid {{
      display: grid;
      grid-template-columns: repeat(auto-fill, minmax(280px, 1fr));
      gap: 1.5rem;
    }}

    .student-card {{
      background: var(--bg-card);
      border: 1px solid var(--border-color);
      border-radius: var(--card-radius);
      padding: 1.25rem;
      backdrop-filter: blur(8px);
      transition: var(--transition);
      display: flex;
      flex-direction: column;
      position: relative;
    }}

    .student-card:hover {{
      transform: translateY(-4px);
      box-shadow: 0 16px 32px -8px rgba(0, 0, 0, 0.5);
      border-color: rgba(56, 189, 248, 0.3);
    }}

    .student-card.flagged {{
      border-color: var(--accent-red);
      background: rgba(239, 68, 68, 0.08);
    }}

    .photo-wrapper {{
      position: relative;
      width: 100%;
      height: 230px;
      border-radius: 12px;
      overflow: hidden;
      background: #090d16;
      display: flex;
      align-items: center;
      justify-content: center;
      cursor: pointer;
      margin-bottom: 1rem;
    }}

    .photo-wrapper img {{
      width: 100%;
      height: 100%;
      object-fit: cover;
      transition: transform 0.3s ease;
    }}

    .photo-wrapper:hover img {{
      transform: scale(1.05);
    }}

    .no-photo-placeholder {{
      display: flex;
      flex-direction: column;
      align-items: center;
      justify-content: center;
      color: var(--text-muted);
      font-size: 0.85rem;
      text-align: center;
      padding: 1rem;
    }}

    .no-photo-icon {{
      font-size: 3rem;
      margin-bottom: 0.5rem;
      opacity: 0.4;
    }}

    .badge-status {{
      position: absolute;
      top: 0.6rem;
      right: 0.6rem;
      padding: 0.25rem 0.65rem;
      border-radius: 999px;
      font-size: 0.7rem;
      font-weight: 700;
      text-transform: uppercase;
      letter-spacing: 0.04em;
    }}

    .badge-status.has-photo {{
      background: rgba(34, 197, 94, 0.85);
      color: #fff;
      backdrop-filter: blur(4px);
    }}

    .badge-status.no-photo {{
      background: rgba(100, 116, 139, 0.7);
      color: #fff;
      backdrop-filter: blur(4px);
    }}

    .badge-status.flagged {{
      background: var(--accent-red);
      color: #fff;
    }}

    .student-name {{
      font-size: 1.05rem;
      font-weight: 700;
      color: #fff;
      line-height: 1.3;
      margin-bottom: 0.3rem;
    }}

    .student-meta {{
      font-size: 0.8rem;
      color: var(--text-muted);
      margin-bottom: 0.75rem;
      display: flex;
      flex-direction: column;
      gap: 0.2rem;
    }}

    .meta-tag {{
      display: inline-block;
      background: rgba(255, 255, 255, 0.06);
      padding: 0.2rem 0.5rem;
      border-radius: 4px;
      font-size: 0.75rem;
      width: fit-content;
    }}

    .student-actions {{
      margin-top: auto;
      display: flex;
      gap: 0.5rem;
      padding-top: 0.75rem;
      border-top: 1px solid var(--border-color);
    }}

    .btn-flag {{
      flex: 1;
      padding: 0.5rem;
      border-radius: 6px;
      font-size: 0.75rem;
      font-weight: 600;
      cursor: pointer;
      transition: var(--transition);
      background: rgba(255, 255, 255, 0.06);
      border: 1px solid var(--border-color);
      color: var(--text-main);
      display: flex;
      align-items: center;
      justify-content: center;
      gap: 0.3rem;
    }}

    .btn-flag:hover {{
      background: rgba(239, 68, 68, 0.2);
      border-color: var(--accent-red);
      color: #fca5a5;
    }}

    .student-card.flagged .btn-flag {{
      background: var(--accent-red);
      border-color: var(--accent-red);
      color: #fff;
    }}

    .btn-copy {{
      padding: 0.5rem 0.75rem;
      border-radius: 6px;
      font-size: 0.75rem;
      cursor: pointer;
      background: rgba(255, 255, 255, 0.06);
      border: 1px solid var(--border-color);
      color: var(--text-muted);
      transition: var(--transition);
    }}

    .btn-copy:hover {{
      color: #fff;
      border-color: var(--accent-blue);
    }}

    /* Modal */
    .modal {{
      display: none;
      position: fixed;
      inset: 0;
      background: rgba(2, 6, 23, 0.85);
      backdrop-filter: blur(10px);
      z-index: 999;
      align-items: center;
      justify-content: center;
      padding: 2rem;
    }}

    .modal.open {{
      display: flex;
    }}

    .modal-content {{
      background: var(--bg-secondary);
      border: 1px solid var(--border-color);
      border-radius: var(--card-radius);
      max-width: 600px;
      width: 100%;
      overflow: hidden;
      box-shadow: 0 25px 50px -12px rgba(0, 0, 0, 0.7);
      animation: modalSlide 0.25s ease-out;
    }}

    @keyframes modalSlide {{
      from {{ transform: scale(0.92); opacity: 0; }}
      to {{ transform: scale(1); opacity: 1; }}
    }}

    .modal-img-wrapper {{
      width: 100%;
      max-height: 450px;
      background: #000;
      display: flex;
      align-items: center;
      justify-content: center;
      overflow: hidden;
    }}

    .modal-img-wrapper img {{
      max-width: 100%;
      max-height: 450px;
      object-fit: contain;
    }}

    .modal-body {{
      padding: 1.5rem;
    }}

    .modal-header {{
      display: flex;
      justify-content: space-between;
      align-items: flex-start;
      margin-bottom: 1rem;
    }}

    .modal-close {{
      background: none;
      border: none;
      color: var(--text-muted);
      font-size: 1.5rem;
      cursor: pointer;
    }}

    .modal-close:hover {{
      color: #fff;
    }}

    /* Toast notification */
    .toast {{
      position: fixed;
      bottom: 2rem;
      right: 2rem;
      background: var(--bg-secondary);
      border: 1px solid var(--accent-blue);
      color: #fff;
      padding: 0.85rem 1.25rem;
      border-radius: 8px;
      font-size: 0.9rem;
      box-shadow: 0 10px 25px rgba(0,0,0,0.5);
      transform: translateY(100px);
      opacity: 0;
      transition: var(--transition);
      z-index: 1000;
    }}

    .toast.show {{
      transform: translateY(0);
      opacity: 1;
    }}
  </style>
</head>
<body>

  <div class="container">
    <header>
      <div class="institution-badge">Instituto Lic. José Vasconcelos (IJOVA)</div>
      <h1>Auditoría y Galería de Fotos de Perfil</h1>
      <div class="subtitle">
        <span>🌐 Dominio: <strong>{domain}</strong></span>
        <span>🕒 Fecha de Auditoría: <strong>{stats['timestamp']}</strong></span>
        <span>⚙️ Microsoft Entra ID Graph API</span>
      </div>

      <div class="metrics-grid">
        <div class="metric-card">
          <div class="metric-label">Total Alumnos Auditados</div>
          <div class="metric-value">{stats['total_audited']}</div>
        </div>
        <div class="metric-card green">
          <div class="metric-label">Con Foto de Perfil</div>
          <div class="metric-value">{stats['with_photo']}</div>
        </div>
        <div class="metric-card amber">
          <div class="metric-label">Sin Foto (Pendientes)</div>
          <div class="metric-value">{stats['without_photo']}</div>
        </div>
        <div class="metric-card">
          <div class="metric-label">Cobertura de Fotos</div>
          <div class="metric-value">{stats['coverage_pct']}%</div>
        </div>
        <div class="metric-card red">
          <div class="metric-label">Marcadas como Inapropiadas</div>
          <div class="metric-value" id="flaggedCount">0</div>
        </div>
      </div>
    </header>

    <div class="toolbar">
      <div class="search-box">
        <span class="search-icon">🔍</span>
        <input type="text" id="searchInput" placeholder="Buscar por alumno, matrícula o correo @{domain}..." autofocus>
      </div>

      <div class="filters">
        <button class="btn active" id="btnFilterPhotos" onclick="setFilter('photo')">📸 Solo con Foto ({stats['with_photo']})</button>
        <button class="btn" id="btnFilterAll" onclick="setFilter('all')">👥 Todos ({stats['total_audited']})</button>
        <button class="btn" id="btnFilterNoPhoto" onclick="setFilter('nophoto')">⚪ Sin Foto ({stats['without_photo']})</button>
        <button class="btn btn-danger" id="btnFilterFlagged" onclick="setFilter('flagged')">⚠️ Inapropiadas (<span id="flaggedBtnCount">0</span>)</button>
        <button class="btn" onclick="exportFlaggedCSV()">📥 Exportar Inapropiadas (CSV)</button>
      </div>
    </div>

    <div class="gallery-grid" id="galleryGrid">
      <!-- Tarjetas insertadas por JavaScript -->
    </div>
  </div>

  <!-- Modal para Inspección Ampliada -->
  <div class="modal" id="photoModal" onclick="closeModal(event)">
    <div class="modal-content" onclick="event.stopPropagation()">
      <div class="modal-img-wrapper">
        <img id="modalImg" src="" alt="Foto ampliada">
      </div>
      <div class="modal-body">
        <div class="modal-header">
          <div>
            <h3 id="modalName" style="color: #fff; font-size: 1.2rem; margin-bottom: 0.25rem;">Nombre del Alumno</h3>
            <p id="modalMeta" style="color: var(--text-muted); font-size: 0.85rem;">Matrícula | Correo</p>
          </div>
          <button class="modal-close" onclick="closeModal()">&times;</button>
        </div>
        <div style="display: flex; gap: 0.5rem; justify-content: flex-end;">
          <button class="btn btn-danger" id="modalFlagBtn" onclick="toggleModalFlag()">⚠️ Marcar Inapropiada</button>
          <button class="btn" onclick="closeModal()">Cerrar</button>
        </div>
      </div>
    </div>
  </div>

  <div class="toast" id="toast">Acción completada</div>

  <script>
    const STUDENTS = {json.dumps(cards_json, ensure_ascii=False)};
    let currentFilter = 'photo';
    let searchQuery = '';
    let flaggedUPNs = new Set(JSON.parse(localStorage.getItem('ijova_flagged_photos') || '[]'));
    let currentModalStudent = null;

    function saveFlags() {{
      localStorage.setItem('ijova_flagged_photos', JSON.stringify(Array.from(flaggedUPNs)));
      updateFlagCounts();
    }}

    function updateFlagCounts() {{
      const count = flaggedUPNs.size;
      document.getElementById('flaggedCount').innerText = count;
      document.getElementById('flaggedBtnCount').innerText = count;
    }}

    function showToast(msg) {{
      const t = document.getElementById('toast');
      t.innerText = msg;
      t.classList.add('show');
      setTimeout(() => t.classList.remove('show'), 2500);
    }}

    function setFilter(filter) {{
      currentFilter = filter;
      document.querySelectorAll('.filters .btn').forEach(b => b.classList.remove('active'));
      if (filter === 'photo') document.getElementById('btnFilterPhotos').classList.add('active');
      if (filter === 'all') document.getElementById('btnFilterAll').classList.add('active');
      if (filter === 'nophoto') document.getElementById('btnFilterNoPhoto').classList.add('active');
      if (filter === 'flagged') document.getElementById('btnFilterFlagged').classList.add('active');
      render();
    }}

    document.getElementById('searchInput').addEventListener('input', (e) => {{
      searchQuery = e.target.value.toLowerCase().trim();
      render();
    }});

    function toggleFlag(upn, e) {{
      if (e) e.stopPropagation();
      if (flaggedUPNs.has(upn)) {{
        flaggedUPNs.delete(upn);
        showToast('Foto desmarcada.');
      }} else {{
        flaggedUPNs.add(upn);
        showToast('⚠️ Alumno marcado como foto inapropiada.');
      }}
      saveFlags();
      render();
    }}

    function copyUPN(upn, e) {{
      if (e) e.stopPropagation();
      navigator.clipboard.writeText(upn).then(() => {{
        showToast('Copiado: ' + upn);
      }});
    }}

    function openModal(student) {{
      if (!student.has_photo) return;
      currentModalStudent = student;
      document.getElementById('modalImg').src = student.img_src;
      document.getElementById('modalName').innerText = student.name;
      document.getElementById('modalMeta').innerText = `${{student.matricula}} | ${{student.upn}} | ${{student.nivel}} (${{student.size_kb}} KB)`;
      
      const flagBtn = document.getElementById('modalFlagBtn');
      if (flaggedUPNs.has(student.upn)) {{
        flagBtn.innerText = '✅ Quitar Marca Inapropiada';
        flagBtn.classList.remove('btn-danger');
      }} else {{
        flagBtn.innerText = '⚠️ Marcar como Inapropiada';
        flagBtn.classList.add('btn-danger');
      }}
      
      document.getElementById('photoModal').classList.add('open');
    }}

    function closeModal() {{
      document.getElementById('photoModal').classList.remove('open');
      currentModalStudent = null;
    }}

    function toggleModalFlag() {{
      if (currentModalStudent) {{
        toggleFlag(currentModalStudent.upn);
        openModal(currentModalStudent);
      }}
    }}

    function exportFlaggedCSV() {{
      if (flaggedUPNs.size === 0) {{
        alert('No has marcado ninguna fotografía como inapropiada.');
        return;
      }}
      const flaggedStudents = STUDENTS.filter(s => flaggedUPNs.has(s.upn));
      let csvContent = 'data:text/csv;charset=utf-8,Matricula,Nombre,Correo,Nivel,Grado,Archivo_Foto\\n';
      flaggedStudents.forEach(s => {{
        csvContent += `"${{s.matricula}}","${{s.name}}","${{s.upn}}","${{s.nivel}}","${{s.grado}}","${{s.filename}}"\\n`;
      }});
      const encodedUri = encodeURI(csvContent);
      const link = document.createElement('a');
      link.setAttribute('href', encodedUri);
      link.setAttribute('download', 'alumnos_fotos_inapropiadas_IJOVA.csv');
      document.body.appendChild(link);
      link.click();
      document.body.removeChild(link);
    }}

    function render() {{
      const grid = document.getElementById('galleryGrid');
      grid.innerHTML = '';

      const filtered = STUDENTS.filter(s => {{
        const isFlagged = flaggedUPNs.has(s.upn);

        // Filter by state
        if (currentFilter === 'photo' && !s.has_photo) return false;
        if (currentFilter === 'nophoto' && s.has_photo) return false;
        if (currentFilter === 'flagged' && !isFlagged) return false;

        // Search query
        if (searchQuery) {{
          const target = `${{s.name}} ${{s.matricula}} ${{s.upn}} ${{s.nivel}}`.toLowerCase();
          if (!target.includes(searchQuery)) return false;
        }}

        return true;
      }});

      if (filtered.length === 0) {{
        grid.innerHTML = `
          <div style="grid-column: 1/-1; text-align: center; padding: 4rem; color: var(--text-muted);">
            <div style="font-size: 3rem; margin-bottom: 1rem;">🔍</div>
            <h3>No se encontraron alumnos con los criterios seleccionados</h3>
            <p style="margin-top: 0.5rem; font-size: 0.9rem;">Prueba cambiando el filtro o la búsqueda.</p>
          </div>
        `;
        return;
      }}

      filtered.forEach(s => {{
        const isFlagged = flaggedUPNs.has(s.upn);
        const card = document.createElement('div');
        card.className = 'student-card' + (isFlagged ? ' flagged' : '');

        let photoHtml = '';
        if (s.has_photo) {{
          photoHtml = `
            <div class="photo-wrapper" onclick='openModal(${{JSON.stringify(s)}})'>
              <img src="${{s.img_src}}" alt="${{s.name}}" loading="lazy">
              <span class="badge-status ${{isFlagged ? 'flagged' : 'has-photo'}}">
                ${{isFlagged ? '⚠️ Inapropiada' : '📸 Foto (' + s.size_kb + ' KB)'}}
              </span>
            </div>
          `;
        }} else {{
          photoHtml = `
            <div class="photo-wrapper" style="cursor: default;">
              <div class="no-photo-placeholder">
                <div class="no-photo-icon">👤</div>
                <div>Sin foto de perfil</div>
              </div>
              <span class="badge-status no-photo">Sin Foto</span>
            </div>
          `;
        }}

        card.innerHTML = `
          ${{photoHtml}}
          <div class="student-name">${{s.name}}</div>
          <div class="student-meta">
            <span><strong>Matrícula:</strong> ${{s.matricula}}</span>
            <span><strong>Correo:</strong> ${{s.upn}}</span>
            <div style="margin-top: 0.3rem;">
              <span class="meta-tag">${{s.nivel}} ${{s.grado}}</span>
            </div>
          </div>
          <div class="student-actions">
            ${{s.has_photo ? `
              <button class="btn-flag" onclick="toggleFlag('${{s.upn}}', event)">
                ${{isFlagged ? '✅ Desmarcar' : '⚠️ Marcar Inapropiada'}}
              </button>
            ` : ''}}
            <button class="btn-copy" title="Copiar correo" onclick="copyUPN('${{s.upn}}', event)">📋 Copiar</button>
          </div>
        `;

        grid.appendChild(card);
      }});
    }}

    // Init
    updateFlagCounts();
    render();
  </script>
</body>
</html>
"""
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(html_content)

    return output_path


def generate_excel_audit_report(
    records: List[PhotoAuditRecord],
    output_path: str,
    stats: Dict[str, Any]
) -> str:
    """
    Genera un archivo Excel (.xlsx) formateado para el expediente de auditoría.
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Auditoría de Fotos"

    # Estilos
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    
    green_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    green_font = Font(name="Calibri", size=10, bold=True, color="166534")
    
    gray_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    gray_font = Font(name="Calibri", size=10, color="64748B")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # Título institucional
    ws.merge_cells("A1:H1")
    ws["A1"] = "INSTITUTO LIC. JOSÉ VASCONCELOS (IJOVA) — AUDITORÍA DE FOTOS DE PERFIL M365"
    ws["A1"].font = Font(name="Calibri", size=13, bold=True, color="0F172A")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # Resumen
    ws["A2"] = f"Fecha: {stats.get('timestamp')} | Total Alumnos: {stats.get('total_audited')} | Con Foto: {stats.get('with_photo')} | Cobertura: {stats.get('coverage_pct')}%"
    ws["A2"].font = Font(name="Calibri", size=10, italic=True, color="475569")
    ws.merge_cells("A2:H2")
    ws.row_dimensions[2].height = 20

    # Encabezados
    headers = [
        "Matrícula",
        "Nombre Completo",
        "Correo Institucional (UPN)",
        "Nivel Escolar",
        "Grado / Semestre",
        "¿Tiene Foto?",
        "Tamaño (KB)",
        "Archivo Guardado"
    ]
    
    ws.row_dimensions[4].height = 24
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # Filas de datos
    for row_idx, r in enumerate(records, start=5):
        ws.row_dimensions[row_idx].height = 20
        size_kb = round(r.photo_size_bytes / 1024, 1) if r.has_photo else 0
        photo_status = "SÍ (CON FOTO)" if r.has_photo else "NO (SIN FOTO)"

        row_data = [
            r.matricula,
            r.display_name,
            r.upn,
            r.nivel,
            r.grado_semestre,
            photo_status,
            size_kb if r.has_photo else "-",
            r.photo_filename or "N/A"
        ]

        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

            # Formateo condicional en columna '¿Tiene Foto?'
            if col_idx == 6:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if r.has_photo:
                    cell.fill = green_fill
                    cell.font = green_font
                else:
                    cell.fill = gray_fill
                    cell.font = gray_font
            elif col_idx in [1, 7]:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    # Ajuste automático de ancho de columnas
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    wb.save(output_path)
    return output_path


def generate_csv_audit_report(
    records: List[PhotoAuditRecord],
    output_path: str
) -> str:
    """
    Genera un archivo CSV simple para procesamiento rápido o importación.
    """
    with open(output_path, "w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow([
            "matricula",
            "nombre_completo",
            "upn",
            "nivel",
            "grado_semestre",
            "tiene_foto",
            "tamano_bytes",
            "archivo_foto"
        ])
        for r in records:
            writer.writerow([
                r.matricula,
                r.display_name,
                r.upn,
                r.nivel,
                r.grado_semestre,
                "SI" if r.has_photo else "NO",
                r.photo_size_bytes,
                r.photo_filename or ""
            ])
    return output_path
