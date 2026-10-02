"""
Pruebas unitarias para el Verificador Global de Cobertura y el Normalizador de Nomenclatura IJOVA.
Verifica:
1. Formato y parseo de nombres segun la formula institucional: [Materia] ([Grado] [Nivel][ Grupo]) - [Ciclo]
2. Deteccion de discrepancias de nomenclatura
3. Auditoria de cobertura global y calculo de porcentajes
4. Generacion de reporte Excel de cobertura
"""

import io
import os
import tempfile
import unittest
from unittest.mock import MagicMock
import openpyxl

from src.teams_engine import (
    format_standard_team_name,
    parse_and_standardize_team_name,
    audit_teams_nomenclature,
    batch_rename_teams,
    audit_global_student_coverage,
    sync_global_student_coverage,
    export_global_coverage_excel,
)


class TestTeamsCoverageAndNomenclature(unittest.TestCase):

    def test_format_standard_team_name(self):
        # Order (subject, grado, nivel, grupo, cycle)
        res1 = format_standard_team_name("Matematicas", "3°", "Secundaria", None, "26-27")
        self.assertEqual(res1, "Matematicas (3° Secundaria) - 26-27")

        # Inverted order (subject, nivel, grado) auto-detected
        res1_inv = format_standard_team_name("Matematicas", "Secundaria", "3°", None, "26-27")
        self.assertEqual(res1_inv, "Matematicas (3° Secundaria) - 26-27")

        # With group
        res2 = format_standard_team_name("Espanol", "1°", "Primaria", "A", "26-27")
        self.assertEqual(res2, "Espanol (1° Primaria A) - 26-27")

    def test_parse_and_standardize_various_names(self):
        test_cases = [
            ("ESPAÑOL III SECUNDARIA 26-27", "Español (3° Secundaria) - 26-27", False),
            ("Ingles 3ro Secundaria", "Inglés (3° Secundaria) - 26-27", False),
            ("EDUCACIÓN FÍSICA II SECUNDARIA 26-27", "Educación Física (2° Secundaria) - 26-27", False),
            ("Español 1 Primaria A", "Español (1° Primaria A) - 26-27", False),
            ("Matemáticas (3° Secundaria) - 26-27", "Matemáticas (3° Secundaria) - 26-27", True),
        ]
        for raw, expected_std, expected_compliant in test_cases:
            parsed = parse_and_standardize_team_name(raw)
            self.assertEqual(parsed["suggested_name"], expected_std)
            self.assertEqual(parsed["is_compliant"], expected_compliant)

    def test_audit_teams_nomenclature(self):
        mock_graph = MagicMock()
        mock_graph.get_all_teams.return_value = [
            {"id": "t1", "displayName": "Matemáticas (3° Secundaria) - 26-27", "description": "Clase"},
            {"id": "t2", "displayName": "Ingles 1ro Secundaria 26-27", "description": "Clase"},
        ]
        res = audit_teams_nomenclature(graph=mock_graph, cycle_filter="2026-2027")
        self.assertTrue(res["success"])
        classes = res["classes"]
        self.assertEqual(len(classes), 2)
        self.assertEqual(res["total_teams"], 2)
        self.assertEqual(res["compliant_count"], 1)
        self.assertEqual(res["non_compliant_count"], 1)

    def test_audit_global_student_coverage(self):
        mock_graph = MagicMock()
        mock_graph.get_all_teams.return_value = [
            {"id": "team-sec-3-esp", "displayName": "Español (3° Secundaria) - 26-27", "description": ""},
            {"id": "team-sec-3-mat", "displayName": "Matemáticas (3° Secundaria) - 26-27", "description": ""},
        ]
        # Student 250001 in both, student 250002 only in esp
        mock_graph.get_team_members.side_effect = lambda tid: [
            {"userPrincipalName": "250001@colegioijova.edu.mx"},
        ] if tid == "team-sec-3-mat" else [
            {"userPrincipalName": "250001@colegioijova.edu.mx"},
            {"userPrincipalName": "250002@colegioijova.edu.mx"},
        ]

        mock_db = {
            "250001": {
                "nombre": "Alumno Uno",
                "grado": "3ro",
                "seccion": "Secundaria",
                "correo": "250001@colegioijova.edu.mx",
                "estatus": "Activo",
            },
            "250002": {
                "nombre": "Alumno Dos",
                "grado": "3ro",
                "seccion": "Secundaria",
                "correo": "250002@colegioijova.edu.mx",
                "estatus": "Activo",
            }
        }

        res = audit_global_student_coverage(graph=mock_graph, school_db=mock_db, cycle_filter="2026-2027")
        self.assertTrue(res["success"])
        students = res["students"]
        self.assertEqual(len(students), 2)

        st1 = next(s for s in students if s["matricula"] == "250001")
        self.assertEqual(st1["coverage_status"], "COBERTURA_COMPLETA")
        self.assertEqual(st1["coverage_percentage"], 100.0)

        st2 = next(s for s in students if s["matricula"] == "250002")
        self.assertEqual(st2["coverage_status"], "COBERTURA_PARCIAL")
        self.assertEqual(st2["coverage_percentage"], 50.0)
        self.assertEqual(len(st2["missing_classes"]), 1)
        self.assertEqual(st2["missing_classes"][0]["subject"], "Matemáticas")

    def test_export_global_coverage_excel(self):
        mock_data = {
            "students": [
                {
                    "matricula": "250019",
                    "display_name": "DIEGO DANIEL GAMALLO",
                    "upn": "250019@colegioijova.edu.mx",
                    "nivel": "Secundaria",
                    "grado": "3°",
                    "total_expected": 6,
                    "enrolled_count": 5,
                    "missing_count": 1,
                    "extraneous_count": 0,
                    "coverage_pct": 83.3,
                    "coverage_status": "INCOMPLETO",
                    "missing_classes": [{"subject": "Química", "team_name": "Química (3° Secundaria) - 26-27"}],
                }
            ],
            "classes": [
                {
                    "team_id": "t1",
                    "name": "Química (3° Secundaria) - 26-27",
                    "target_grade": "3° Secundaria",
                    "members_count": 28,
                    "expected_count": 29,
                    "missing_count": 1,
                    "extraneous_count": 0,
                    "balance_status": "Faltan Alumnos",
                }
            ],
            "summary": {
                "total_students_audited": 1,
                "full_coverage_students": 0,
                "unmet_coverage_students": 1,
                "active_classes_evaluated": 1,
                "global_coverage_rate": 83.3,
            }
        }
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tf:
            tmp_path = tf.name

        try:
            out_file = export_global_coverage_excel(mock_data, tmp_path)
            self.assertTrue(os.path.exists(out_file))
            wb = openpyxl.load_workbook(out_file)
            self.assertIn("Cobertura Alumnos", wb.sheetnames)
            self.assertIn("Estado de Clases", wb.sheetnames)
            ws1 = wb["Cobertura Alumnos"]
            self.assertEqual(ws1.cell(row=2, column=1).value, "250019")
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)


if __name__ == "__main__":
    unittest.main()
