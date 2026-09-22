"""
Pruebas unitarias para el motor de auditoría de tareas escolares y cumplimiento docente.
Verifica auditoría individual por clase, auditoría masiva por ciclo escolar,
cálculo de semáforo de actividad docente y generación de reporte ejecutivo Excel.
"""
import os
import shutil
import tempfile
import unittest
from unittest.mock import MagicMock
import openpyxl

from src.teams_engine import (
    audit_class_assignments,
    audit_all_assignments,
    export_assignments_report_excel,
)


class TestAssignmentsEngine(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_audit_class_assignments(self):
        """Verifica la auditoría de tareas y entregas de una clase individual."""
        mock_graph = MagicMock()

        # Mock de 2 tareas en la clase
        mock_graph.get_class_assignments.return_value = [
            {
                "id": "task-001",
                "displayName": "Ensayo Histórico",
                "status": "assigned",
                "dueDateTime": "2026-09-30T23:59:00Z",
                "grading": {"maxPoints": 100},
                "instructions": {"content": "Escribir un ensayo de 2 cuartillas."}
            },
            {
                "id": "task-002",
                "displayName": "Cuestionario de Comprensión",
                "status": "assigned",
                "dueDateTime": "2026-10-05T23:59:00Z",
                "grading": None,
                "instructions": {"content": "Responder preguntas 1 a 10."}
            }
        ]

        # Mock de entregas: tarea 1 tiene 2 entregadas y 1 pendiente
        def mock_submissions(class_id, assignment_id):
            if assignment_id == "task-001":
                return [
                    {"id": "sub-1", "status": "submitted"},
                    {"id": "sub-2", "status": "turnedIn"},
                    {"id": "sub-3", "status": "working"},
                ]
            else:
                return [
                    {"id": "sub-4", "status": "submitted"},
                    {"id": "sub-5", "status": "working"},
                ]

        mock_graph.get_assignment_submissions.side_effect = mock_submissions

        result = audit_class_assignments(mock_graph, "class-123")

        self.assertEqual(result["class_id"], "class-123")
        self.assertEqual(result["total_assignments"], 2)
        self.assertEqual(result["total_submissions"], 5)
        self.assertEqual(result["total_turned_in"], 3)
        self.assertEqual(result["turn_in_rate"], 60.0)

        # Verificar desglose de la primera tarea
        t1 = result["assignments"][0]
        self.assertEqual(t1["id"], "task-001")
        self.assertEqual(t1["title"], "Ensayo Histórico")
        self.assertEqual(t1["points"], 100)
        self.assertEqual(t1["submissions_count"], 3)
        self.assertEqual(t1["turned_in_count"], 2)
        self.assertEqual(t1["turn_in_rate"], 66.7)

    def test_audit_all_assignments(self):
        """Verifica la auditoría consolidada por ciclo y el cálculo de semáforo docente."""
        mock_graph = MagicMock()

        # Mock de teams_list: 2 clases del ciclo 2026-2027 y 1 del ciclo 2025-2026
        mock_teams_data = {
            "teams": [
                {
                    "id": "c1",
                    "name": "Matemáticas 3er Semestre - 26-27",
                    "team_type": "CLASE",
                    "academic_cycle": "2026-2027",
                    "owners": [{"name": "Profesor Gómez", "userPrincipalName": "jgomez@ijova.com"}],
                    "students_count": 25,
                },
                {
                    "id": "c2",
                    "name": "Historia 1° Secundaria - 26-27",
                    "team_type": "CLASE",
                    "academic_cycle": "2026-2027",
                    "owners": [{"name": "Profesora Morales", "userPrincipalName": "lmorales@ijova.com"}],
                    "students_count": 20,
                },
                {
                    "id": "c3",
                    "name": "Física Anterior - 25-26",
                    "team_type": "CLASE",
                    "academic_cycle": "2025-2026",
                    "owners": [{"name": "Profesor Gómez", "userPrincipalName": "jgomez@ijova.com"}],
                    "students_count": 15,
                }
            ]
        }

        # c1 tiene 4 tareas (docente activo >= 4)
        def mock_assignments(class_id):
            if class_id == "c1":
                return [
                    {"id": f"t-{i}", "displayName": f"Tarea {i}", "status": "assigned", "dueDateTime": None, "grading": None, "instructions": None}
                    for i in range(1, 5)
                ]
            elif class_id == "c2":
                return []  # 0 tareas (docente inactivo)
            return []

        mock_graph.get_class_assignments.side_effect = mock_assignments
        mock_graph.get_assignment_submissions.return_value = []

        result = audit_all_assignments(mock_graph, target_cycle="2026-2027", teams_data=mock_teams_data)

        summary = result["summary"]
        self.assertEqual(summary["cycle"], "2026-2027")
        self.assertEqual(summary["total_classes"], 2)
        self.assertEqual(summary["classes_with_assignments"], 1)
        self.assertEqual(summary["classes_without_assignments"], 1)
        self.assertEqual(summary["total_assignments"], 4)
        self.assertEqual(summary["docentes_activos"], 1)
        self.assertEqual(summary["docentes_inactivos"], 1)
        self.assertEqual(summary["docentes_moderados"], 0)

        # Verificar registro de docentes
        docentes = result["docentes"]
        self.assertEqual(len(docentes), 2)
        gomez = next((d for d in docentes if "jgomez" in d["upn"]), None)
        self.assertIsNotNone(gomez)
        self.assertEqual(gomez["total_assignments"], 4)
        self.assertEqual(gomez["status"], "ACTIVO")
        self.assertIn("Activo", gomez["status_tag"])

    def test_export_assignments_report_excel(self):
        """Verifica la generación del libro Excel institucional con 3 hojas formales."""
        mock_data = {
            "summary": {
                "cycle": "2026-2027",
                "total_classes": 2,
                "classes_with_assignments": 1,
                "classes_without_assignments": 1,
                "total_assignments": 2,
                "total_submissions": 10,
                "total_turned_in": 7,
                "overall_turn_in_rate": 70.0,
                "docentes_activos": 0,
                "docentes_moderados": 1,
                "docentes_inactivos": 1,
            },
            "docentes": [
                {
                    "name": "Profesor Gómez",
                    "upn": "jgomez@ijova.com",
                    "total_classes": 1,
                    "classes_with_tasks": 1,
                    "total_assignments": 2,
                    "total_submissions": 10,
                    "total_turned_in": 7,
                    "turn_in_rate": 70.0,
                    "status_tag": "Moderado (1-3 tareas)",
                },
                {
                    "name": "Profesora Morales",
                    "upn": "lmorales@ijova.com",
                    "total_classes": 1,
                    "classes_with_tasks": 0,
                    "total_assignments": 0,
                    "total_submissions": 0,
                    "total_turned_in": 0,
                    "turn_in_rate": 0.0,
                    "status_tag": "Inactivo (0 tareas)",
                }
            ],
            "classes": [
                {
                    "id": "c1",
                    "name": "Matemáticas 3er Semestre",
                    "academic_cycle": "2026-2027",
                    "teacher_name": "Profesor Gómez",
                    "teacher_upn": "jgomez@ijova.com",
                    "total_assignments": 2,
                    "total_submissions": 10,
                    "total_turned_in": 7,
                    "turn_in_rate": 70.0,
                    "assignments": [
                        {
                            "id": "t1",
                            "title": "Álgebra Lineal",
                            "status": "assigned",
                            "due_date": "2026-09-30T18:00:00Z",
                            "points": 100,
                            "instructions": "Resolver ejercicios 1 al 5.",
                            "submissions_count": 5,
                            "turned_in_count": 4,
                            "turn_in_rate": 80.0,
                        },
                        {
                            "id": "t2",
                            "title": "Matrices y Determinantes",
                            "status": "assigned",
                            "due_date": "2026-10-02T18:00:00Z",
                            "points": 50,
                            "instructions": "Resolver problemario.",
                            "submissions_count": 5,
                            "turned_in_count": 3,
                            "turn_in_rate": 60.0,
                        }
                    ]
                }
            ]
        }

        output_path = os.path.join(self.temp_dir, "Reporte_Tareas_Test.xlsx")
        export_assignments_report_excel(mock_data, output_path)

        self.assertTrue(os.path.exists(output_path))
        self.assertGreater(os.path.getsize(output_path), 5000)

        # Inspeccionar estructura del libro
        wb = openpyxl.load_workbook(output_path)
        self.assertIn("Resumen Ejecutivo Tareas", wb.sheetnames)
        self.assertIn("Semaforo Cumplimiento Docente", wb.sheetnames)
        self.assertIn("Bitacora Detallada de Tareas", wb.sheetnames)

        # Verificar hoja de resumen
        ws_summary = wb["Resumen Ejecutivo Tareas"]
        self.assertIn("INFORME OFICIAL DE AUDITORÍA", ws_summary["A2"].value)

        # Verificar hoja de docentes
        ws_docentes = wb["Semaforo Cumplimiento Docente"]
        self.assertEqual(ws_docentes["A2"].value, "Profesor Gómez")
        self.assertEqual(ws_docentes["F2"].value, "Moderado (1-3 tareas)")

        # Verificar hoja detallada
        ws_tasks = wb["Bitacora Detallada de Tareas"]
        self.assertEqual(ws_tasks["A2"].value, "Matemáticas 3er Semestre")
        self.assertEqual(ws_tasks["E2"].value, "Álgebra Lineal")


if __name__ == "__main__":
    unittest.main()
