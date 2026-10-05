"""
Pruebas unitarias para el Auditor de Inactividad y Primer Acceso (Teams y M365).
Verifica:
1. Logica de clasificacion: sesiones abiertas (movil/PC), cuentas sin primer acceso (clave temporal pendiente), reseteos recientes y cuentas suspendidas
2. Cruce con la plantilla oficial del ciclo 2026-2027 y base de tutores
3. Generacion de archivo Excel de auditoria con columnas actualizadas
4. Endpoints de la interfaz web Flask
"""

import os
import tempfile
import unittest
from datetime import datetime, timezone, timedelta
from unittest.mock import MagicMock, patch
import openpyxl

from src.teams_engine import (
    audit_students_login_activity,
    export_activity_audit_excel,
)
from src.gui.app import create_app


class TestTeamsActivityAudit(unittest.TestCase):

    def setUp(self):
        self.mock_graph = MagicMock()
        self.mock_db = {
            "250001": {
                "matricula": "250001",
                "nombre_oficial": "ALUMNO SESION RECIENTE",
                "upn": "250001@ijova.com",
                "nivel": "Secundaria",
                "grado_semestre": "3°",
                "estatus": "Activo",
                "padre_o_tutor": "PADRE UNO",
                "telefono_contacto": "5511223344",
                "correo_contacto": "padre1@correo.com",
            },
            "250002": {
                "matricula": "250002",
                "nombre_oficial": "ALUMNO SESION MOVIL PERSISTENTE",
                "upn": "250002@ijova.com",
                "nivel": "Secundaria",
                "grado_semestre": "2°",
                "estatus": "Activo",
                "padre_o_tutor": "MADRE DOS",
                "telefono_contacto": "5522334455",
                "correo_contacto": "madre2@correo.com",
            },
            "250003": {
                "matricula": "250003",
                "nombre_oficial": "ALUMNO SIN INICIO DE SESION",
                "upn": "250003@ijova.com",
                "nivel": "Primaria",
                "grado_semestre": "5°",
                "estatus": "Activo",
                "padre_o_tutor": "TUTOR TRES",
                "telefono_contacto": "5533445566",
                "correo_contacto": "tutor3@correo.com",
            },
            "250004": {
                "matricula": "250004",
                "nombre_oficial": "ALUMNO CUENTA SUSPENDIDA",
                "upn": "250004@ijova.com",
                "nivel": "Preparatoria",
                "grado_semestre": "1°",
                "estatus": "Activo",
                "padre_o_tutor": "TUTOR CUATRO",
                "telefono_contacto": "5544556677",
                "correo_contacto": "tutor4@correo.com",
            },
            "240001": {
                "matricula": "240001",
                "nombre_oficial": "ALUMNO BAJA O CICLO ANTERIOR",
                "upn": "240001@ijova.com",
                "nivel": "Secundaria",
                "grado_semestre": "3°",
                "estatus": "Baja",
            }
        }

    def test_audit_students_login_activity_classification(self):
        now = datetime.now(timezone.utc)

        ts_recent = (now - timedelta(days=2)).strftime("%Y-%m-%dT%H:%M:%SZ")
        ts_older = (now - timedelta(days=31)).strftime("%Y-%m-%dT%H:%M:%SZ")
        ts_created = (now - timedelta(days=40)).strftime("%Y-%m-%dT10:00:00Z")

        graph_users = [
            # 1. Sesión abierta y reseteo reciente (<5d)
            {
                "id": "u1",
                "displayName": "ALUMNO SESION RECIENTE",
                "userPrincipalName": "250001@ijova.com",
                "createdDateTime": ts_created,
                "signInSessionsValidFromDateTime": ts_recent,
                "deviceKeys": [{"deviceId": "d1"}],
                "accountEnabled": True,
                "passwordProfile": None,
            },
            # 2. Sesión abierta en móvil (token persistente, >5d desde último reseteo, password propio)
            {
                "id": "u2",
                "displayName": "ALUMNO SESION MOVIL PERSISTENTE",
                "userPrincipalName": "250002@ijova.com",
                "createdDateTime": ts_created,
                "signInSessionsValidFromDateTime": ts_older,
                "deviceKeys": [],
                "accountEnabled": True,
                "passwordProfile": None,
            },
            # 3. Nunca inició sesión: tiene contraseña temporal pendiente de cambio
            {
                "id": "u3",
                "displayName": "ALUMNO SIN INICIO DE SESION",
                "userPrincipalName": "250003@ijova.com",
                "createdDateTime": ts_created,
                "signInSessionsValidFromDateTime": ts_older,
                "deviceKeys": [],
                "accountEnabled": True,
                "passwordProfile": {"forceChangePasswordNextSignIn": True},
            },
            # 4. Cuenta suspendida / deshabilitada
            {
                "id": "u4",
                "displayName": "ALUMNO CUENTA SUSPENDIDA",
                "userPrincipalName": "250004@ijova.com",
                "createdDateTime": ts_created,
                "signInSessionsValidFromDateTime": ts_older,
                "deviceKeys": [],
                "accountEnabled": False,
                "passwordProfile": None,
            },
            # 5. Excluido por no estar activo en plantilla oficial
            {
                "id": "u5",
                "displayName": "ALUMNO EXCLUIDO",
                "userPrincipalName": "240001@ijova.com",
                "createdDateTime": ts_created,
                "signInSessionsValidFromDateTime": ts_recent,
                "deviceKeys": [],
                "accountEnabled": True,
                "passwordProfile": None,
            }
        ]

        self.mock_graph.get_users_with_activity.return_value = graph_users

        res = audit_students_login_activity(self.mock_graph, self.mock_db, days_threshold=5)

        summary = res["summary"]
        students = res["students"]

        # Exactamente 4 alumnos oficiales activos evaluados (240001 excluido)
        self.assertEqual(summary["total_students"], 4)
        self.assertEqual(summary["active_sessions"], 2)
        self.assertEqual(summary["never_logged_in"], 1)
        self.assertEqual(summary["recent_resets"], 1)
        self.assertEqual(summary["disabled_accounts"], 1)
        self.assertEqual(summary["adoption_rate"], 50.0)

        stu_map = {s["matricula"]: s for s in students}

        # Alumno 1: Sesión abierta, reseteo reciente
        self.assertEqual(stu_map["250001"]["session_status"], "RECENT_RESET")
        self.assertTrue(stu_map["250001"]["has_active_session"])
        self.assertFalse(stu_map["250001"]["never_logged_in"])
        self.assertEqual(stu_map["250001"]["devices_count"], 1)
        self.assertEqual(stu_map["250001"]["tutor"]["nombre"], "PADRE UNO")

        # Alumno 2: Sesión abierta en móvil/dispositivo
        self.assertEqual(stu_map["250002"]["session_status"], "ACTIVE_SESSION")
        self.assertTrue(stu_map["250002"]["has_active_session"])
        self.assertFalse(stu_map["250002"]["never_logged_in"])
        self.assertEqual(stu_map["250002"]["tutor"]["telefono"], "5522334455")

        # Alumno 3: Sin inicio de sesión (Clave provisional pendiente)
        self.assertEqual(stu_map["250003"]["session_status"], "NEVER_LOGGED_IN")
        self.assertFalse(stu_map["250003"]["has_active_session"])
        self.assertTrue(stu_map["250003"]["never_logged_in"])
        self.assertEqual(stu_map["250003"]["risk_level"], "CRITICAL")
        self.assertEqual(stu_map["250003"]["credential_type"], "Clave temporal no utilizada")

        # Alumno 4: Cuenta suspendida
        self.assertEqual(stu_map["250004"]["session_status"], "DISABLED")
        self.assertFalse(stu_map["250004"]["has_active_session"])
        self.assertEqual(stu_map["250004"]["risk_level"], "CRITICAL")

    def test_export_activity_audit_excel(self):
        activity_data = {
            "summary": {
                "total_students": 2,
                "active_sessions": 1,
                "never_logged_in": 1,
                "recent_resets": 0,
                "disabled_accounts": 0,
                "adoption_rate": 50.0,
                "threshold_days": 5,
            },
            "students": [
                {
                    "matricula": "250001",
                    "displayName": "ALUMNO CON SESION",
                    "userPrincipalName": "250001@ijova.com",
                    "nivel": "Secundaria",
                    "grado": "3°",
                    "created_date": "2026-08-20",
                    "last_session_date": "2026-09-04 12:00",
                    "session_status": "ACTIVE_SESSION",
                    "status_text": "Sesión abierta en móvil/dispositivo",
                    "credential_type": "Contraseña personal configurada",
                    "has_active_session": True,
                    "never_logged_in": False,
                    "risk_level": "ACTIVE",
                    "devices_count": 0,
                    "account_enabled": True,
                    "tutor": {
                        "nombre": "TUTOR UNO",
                        "telefono": "5512345678",
                        "correo": "tutor1@mail.com",
                    }
                },
                {
                    "matricula": "250002",
                    "displayName": "ALUMNO SIN SESION",
                    "userPrincipalName": "250002@ijova.com",
                    "nivel": "Secundaria",
                    "grado": "2°",
                    "created_date": "2026-08-20",
                    "last_session_date": "Sin registro",
                    "session_status": "NEVER_LOGGED_IN",
                    "status_text": "Sin inicio de sesión (Clave provisional pendiente)",
                    "credential_type": "Clave temporal no utilizada",
                    "has_active_session": False,
                    "never_logged_in": True,
                    "risk_level": "CRITICAL",
                    "devices_count": 0,
                    "account_enabled": True,
                    "tutor": {
                        "nombre": "TUTOR DOS",
                        "telefono": "5587654321",
                        "correo": "tutor2@mail.com",
                    }
                }
            ]
        }

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            tmp_path = tmp.name

        try:
            res_path = export_activity_audit_excel(activity_data, tmp_path)
            self.assertEqual(res_path, tmp_path)
            self.assertTrue(os.path.exists(tmp_path))

            wb = openpyxl.load_workbook(tmp_path)
            self.assertIn("Auditoría de Inactividad", wb.sheetnames)
            ws = wb["Auditoría de Inactividad"]

            # Fila 1 son encabezados oficiales
            headers = [cell.value for cell in ws[1]]
            expected_headers = [
                "Matrícula",
                "Nombre del Alumno",
                "Nivel",
                "Grado",
                "Correo Institucional",
                "Estado Actividad",
                "Tipo de Credencial",
                "Dispositivos",
                "Última Sesión",
                "Fecha Creación",
                "Nombre Tutor",
                "Teléfono Tutor",
                "Correo Tutor"
            ]
            self.assertEqual(headers, expected_headers)

            # Fila 2 y 3 datos de alumnos
            self.assertEqual(ws.cell(row=2, column=1).value, "250001")
            self.assertEqual(ws.cell(row=2, column=6).value, "Sesión abierta en móvil/dispositivo")
            self.assertEqual(ws.cell(row=2, column=7).value, "Contraseña personal configurada")

            self.assertEqual(ws.cell(row=3, column=1).value, "250002")
            self.assertEqual(ws.cell(row=3, column=6).value, "Sin inicio de sesión (Clave provisional pendiente)")
            self.assertEqual(ws.cell(row=3, column=7).value, "Clave temporal no utilizada")
            wb.close()
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)

    def test_gui_activity_endpoints(self):
        app = create_app()
        client = app.test_client()

        sample_audit_result = {
            "success": True,
            "summary": {
                "total_students": 1,
                "active_sessions": 1,
                "never_logged_in": 0,
                "recent_resets": 0,
                "disabled_accounts": 0,
                "adoption_rate": 100.0,
                "threshold_days": 5,
            },
            "students": [
                {
                    "matricula": "250010",
                    "displayName": "TEST ALUMNO",
                    "userPrincipalName": "250010@ijova.com",
                    "nivel": "Primaria",
                    "grado": "1°",
                    "last_sign_in_formatted": "05/10/2026 08:00",
                    "has_active_session": True,
                    "never_logged_in": False,
                    "session_status": "ACTIVE_SESSION",
                    "status_text": "Sesión abierta en móvil/dispositivo",
                    "credential_type": "Contraseña personal configurada",
                    "risk_level": "ACTIVE",
                    "devices_count": 1,
                    "account_enabled": True,
                    "tutor": {}
                }
            ]
        }

        with patch("src.teams_engine.audit_students_login_activity", return_value=sample_audit_result):
            resp = client.get("/api/teams/activity/audit")
            self.assertEqual(resp.status_code, 200)
            data = resp.get_json()
            self.assertTrue(data["success"])
            self.assertEqual(data["data"]["summary"]["active_sessions"], 1)

        with patch("src.teams_engine.audit_students_login_activity", return_value=sample_audit_result):
            resp_excel = client.get("/api/teams/activity/export-excel")
            self.assertEqual(resp_excel.status_code, 200)
            self.assertIn("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", resp_excel.content_type)
            resp_excel.close()

if __name__ == "__main__":
    unittest.main()
